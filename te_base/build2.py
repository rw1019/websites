"""按《Te 数据表》新规格重构：产出 5 张表 + 导出内嵌图片。"""
import json
import math
import os
import re
import shutil
import xml.etree.ElementTree as ET
from datetime import datetime, date, timedelta

import openpyxl

SRC = "/workspace/.uploads/04716e07-c586-4875-b818-21b23fe18e64_Te 数据表.xlsx"
XDIR = "/workspace/te_base/xlsx"
OUT = "/workspace/te_base/out2"
IMG = "/workspace/te_base/out2/images"
BASE_DATE = date(2026, 10, 19)          # 备货预警固定基准时间（替代 TODAY()）
ZONE = "滞销/其他汇总区"                 # 滞销款归入的汇总主分类

NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def clean_header(s):
    return None if s is None else re.sub(r"\s+", "", str(s)).strip()


def to_date(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)) and 30000 <= v <= 60000:
        try:
            return openpyxl.utils.datetime.from_excel(v).strftime("%Y-%m-%d")
        except Exception:
            return v
    return v


def sanitize_text(v):
    if isinstance(v, str):
        v = re.sub(r"[\x00-\x08\x0b-\x1f\x7f]", " ", v).strip()
        return v or None
    return v


def is_date_str(v):
    return isinstance(v, str) and bool(re.match(r"^\d{4}-\d{2}-\d{2}$", v))


def to_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except Exception:
        return None


# ---------- 内嵌图片 ----------
rels = ET.parse(f"{XDIR}/xl/_rels/cellimages.xml.rels").getroot()
rid2file = {rel.get("Id"): os.path.basename(rel.get("Target")) for rel in rels}
cellimg = ET.parse(f"{XDIR}/xl/cellimages.xml").getroot()
id2file = {}
for ci in cellimg:
    pic = ci.find("xdr:pic", NS)
    if pic is None:
        continue
    name = pic.find("xdr:nvPicPr", NS).find("xdr:cNvPr", NS).get("name")
    rid = pic.find("xdr:blipFill", NS).find("a:blip", NS).get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
    id2file[name] = rid2file.get(rid)


def dispimg_id(cellval):
    if isinstance(cellval, str) and "DISPIMG" in cellval:
        m = re.search(r'DISPIMG\("([^"]+)"', cellval)
        return m.group(1) if m else None
    return None


os.makedirs(IMG, exist_ok=True)
os.makedirs(OUT, exist_ok=True)
for f in os.listdir(IMG):
    os.remove(os.path.join(IMG, f))

wb = openpyxl.load_workbook(SRC, data_only=True)
tables = {}

# ================= 表1 采购发货流水 / 表2 SKU档案库 =================
ws = wb["库存"]
heads = [hr for rng in ws.merged_cells.ranges if str(rng).startswith("A")
         for hr in [rng.min_row] if isinstance(ws.cell(rng.min_row, 1).value, str)]
heads.sort()
bounds = heads + [ws.max_row + 1]

blocks = []
for i, hr in enumerate(heads):
    blocks.append({"hr": hr, "end": bounds[i + 1],
                   "raw": str(ws.cell(hr, 1).value).strip(),
                   "img_id": dispimg_id(ws.cell(hr, 15).value),
                   "skc": sanitize_text(ws.cell(hr, 16).value)})

parent = None
for b in blocks:
    if b["img_id"]:                        # 有图 = 真产品
        parent = b
        b["主分类"], b["子规格"] = b["raw"], None
    else:                                  # 无图 = 上级产品的颜色/规格
        b["主分类"], b["子规格"] = (parent["raw"] if parent else b["raw"]), b["raw"]
# 滞销款单独归入汇总区
for b in blocks:
    if b["raw"] == "滞销款":
        b["主分类"], b["子规格"] = ZONE, "滞销款"
for b in blocks:
    b["完整名称"] = f"{b['主分类']}-{b['子规格']}" if b["子规格"] else b["主分类"]

flows = []
for b in blocks:
    for r in range(b["hr"] + 1, b["end"]):
        if ws.cell(r, 4).value is None and ws.cell(r, 9).value is None:
            continue
        pdate, sdate = to_date(ws.cell(r, 1).value), to_date(ws.cell(r, 2).value)
        notes = []
        n0 = sanitize_text(ws.cell(r, 17).value)
        if n0 is not None:
            notes.append(str(n0))
        for label, d in (("采购", pdate), ("发货", sdate)):
            if d is not None and not is_date_str(d):
                notes.append(f"{label}日期备注：{d}")
                pdate = None if label == "采购" else pdate
                sdate = None if label == "发货" else sdate
        qty = to_num(ws.cell(r, 9).value)
        price = to_num(ws.cell(r, 4).value)
        freight = to_num(ws.cell(r, 5).value)
        temu = to_num(ws.cell(r, 6).value)
        unit_cost = to_num(ws.cell(r, 8).value)
        if unit_cost is None and price is not None and qty:
            unit_cost = price + (freight or 0) / qty + (temu or 0) / qty
        total = to_num(ws.cell(r, 13).value)
        if total is None and price is not None:
            total = price * (qty or 0) + (freight or 0) + (temu or 0)
        declared = to_num(ws.cell(r, 10).value)
        profit = to_num(ws.cell(r, 12).value)
        flows.append({
            "主分类": b["主分类"],
            "子规格": b["子规格"],
            "产品完整名称": b["完整名称"],
            "采购日期": pdate,
            "采购年月": pdate[:7] if is_date_str(pdate) else None,
            "发货日期": sdate,
            "发货状态": "已发货" if is_date_str(sdate) else "待发货",
            "采购单价": price,
            "采购数量": qty,
            "国内运费": freight,
            "TEMU运费": temu,
            "单个综合成本": unit_cost,
            "本次总支出": total,
            "申报价": declared,
            "单件利润": profit,
            "净利率": to_num(ws.cell(r, 11).value),
            "预计销售总额": declared * qty if declared is not None and qty is not None else None,
            "预计毛利总额": profit * qty if profit is not None and qty is not None else None,
            "SKC": b["skc"],
            "备注": "；".join(n for n in notes if n) or None,
        })
tables["采购发货流水"] = flows

# 表2：按 完整名称 聚合
sku = {}
for b in blocks:
    sku[b["完整名称"]] = {"产品完整名称": b["完整名称"], "所属大类": b["主分类"],
                          "SKC编码": b["skc"], "图片关联ID": b["img_id"], "_img": None,
                          "_rows": []}
for f in flows:
    sku[f["产品完整名称"]]["_rows"].append(f)
for name, s in sku.items():
    rs = s.pop("_rows")
    prices = [(r["采购日期"], r["采购单价"]) for r in rs if r["采购单价"] is not None]
    decls = [(r["采购日期"], r["申报价"]) for r in rs if r["申报价"] is not None]
    s["最新/常用采购价"] = max(prices, key=lambda x: x[0] or "")[1] if prices else None
    s["最新申报价"] = max(decls, key=lambda x: x[0] or "")[1] if decls else None
    s["累计采购批次"] = len(rs)
    s["累计采购总量"] = round(sum(r["采购数量"] or 0 for r in rs), 2)
    s["累计采购总金额"] = round(sum(r["本次总支出"] or 0 for r in rs), 2)
    s["累计理论毛利贡献"] = round(sum(r["预计毛利总额"] or 0 for r in rs), 2)
tables["SKU档案库"] = list(sku.values())

# 导出图片：以 完整名称 命名
for b in blocks:
    if b["img_id"] and id2file.get(b["img_id"]):
        shutil.copy(os.path.join(XDIR, "xl/media", id2file[b["img_id"]]),
                    os.path.join(IMG, f"{b['完整名称']}.png"))
        for s in tables["SKU档案库"]:
            if s["产品完整名称"] == b["完整名称"]:
                s["_img"] = f"{b['完整名称']}.png"

# ================= 表3 物料采购流水 =================
ws = wb["物料"]
mat = []
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 2).value is None and ws.cell(r, 3).value is None:
        continue
    d = to_date(ws.cell(r, 1).value)
    mat.append({"采购日期": d, "采购月份": d[:7] if is_date_str(d) else None,
                "耗材品名": sanitize_text(ws.cell(r, 2).value),
                "支出金额": to_num(ws.cell(r, 3).value)})
tables["物料采购流水"] = mat

# ================= 表4 月度资金对账 =================
inv_m, mat_m = {}, {}
for f in flows:
    if f["采购年月"]:
        inv_m[f["采购年月"]] = inv_m.get(f["采购年月"], 0) + (f["本次总支出"] or 0)
for m in mat:
    if m["采购月份"]:
        mat_m[m["采购月份"]] = mat_m.get(m["采购月份"], 0) + (m["支出金额"] or 0)

ws = wb["物料"]
recon, roll = [], 0.0
for r in range(3, 30):
    d = to_date(ws.cell(r, 71).value)
    if r == 4:                       # 源表把 2024-08 误写成 2025-08（序列号差一年）
        d = "2024-08-31"
    m = d[:7]
    pay = to_num(ws.cell(r, 72).value) or 0        # 平台结算回款
    sub = to_num(ws.cell(r, 73).value) or 0        # 其他收入/补贴
    spend = round(inv_m.get(m, 0) + mat_m.get(m, 0), 2)   # 当月总支出=表1货款+表3物料费
    net = round(pay + sub - spend, 2)
    roll = round(roll + net, 2)
    recon.append({"对账月份": d, "平台结算回款": pay or None, "其他收入/补贴": sub or None,
                  "当月总支出": spend, "当月现金净结余": net, "累计滚存资金": roll})
tables["月度资金对账"] = recon

# ================= 表5a Y2 半托管测算 =================
ws = wb["Y2"]
y2, seen = [], set()
for r in range(2, ws.max_row + 1):
    name = sanitize_text(ws.cell(r, 2).value)
    if not name or name in seen:
        continue
    seen.add(name)
    cost = to_num(ws.cell(r, 12).value)
    decl = to_num(ws.cell(r, 4).value)
    y2.append({
        "产品名": name, "订单日期": to_date(ws.cell(r, 1).value),
        "采购价": to_num(ws.cell(r, 3).value), "申报价": decl, "活动价": to_num(ws.cell(r, 5).value),
        "重量kg": to_num(ws.cell(r, 6).value), "国内快递": to_num(ws.cell(r, 7).value),
        "Y2头程": to_num(ws.cell(r, 8).value), "贴单费": to_num(ws.cell(r, 9).value),
        "尾程运费": to_num(ws.cell(r, 10).value), "汇率": 6.8,
        "成本价": cost, "每单赚": to_num(ws.cell(r, 13).value),
        "毛利率": round((decl - cost) / decl, 4) if decl and cost is not None else None,
    })
tables["Y2测算"] = y2

# ================= 表5b 备货预警 =================
ws = wb["春节备货"]
base_serial = (BASE_DATE - date(1899, 12, 30)).days
stock = []
for r in range(2, ws.max_row + 1):
    name = sanitize_text(ws.cell(r, 1).value)
    if not name:
        continue
    b7, b30 = to_num(ws.cell(r, 2).value), to_num(ws.cell(r, 3).value)
    avail = sum((to_num(ws.cell(r, c).value) or 0) for c in (4, 5, 6))
    lead = sum((to_num(ws.cell(r, c).value) or 0) for c in (7, 8, 9))
    jserial = to_num(ws.cell(r, 10).value)
    d7 = (b7 or 0) / 7
    d30 = (b30 or 0) / 30
    k7 = (jserial - base_serial + lead) * d7 - avail
    k30 = (jserial - base_serial + lead) * d30 - avail
    price = to_num(ws.cell(r, 13).value)
    stock.append({
        "产品名称": name, "最近7天销量": b7, "最近30天销量": b30,
        "7天日均销速": round(d7, 4), "30天日均销速": round(d30, 4),
        "仓内可用": to_num(ws.cell(r, 4).value), "已发库存": to_num(ws.cell(r, 5).value),
        "已采购": to_num(ws.cell(r, 6).value), "在途可用库存": avail,
        "卖家发货天数": to_num(ws.cell(r, 7).value), "物流天数": to_num(ws.cell(r, 8).value),
        "入库时长": to_num(ws.cell(r, 9).value),
        "预计可销售到": to_date(jserial),
        "安全采购量(7天)": round(k7, 2), "安全采购量(30天)": round(k30, 2),
        "单价": price, "建议采购总金额": round((k7 + k30) / 2 * (price or 0), 2),
        "采购状态": sanitize_text(ws.cell(r, 15).value),
    })
tables["备货预警"] = stock

with open(f"{OUT}/tables.json", "w", encoding="utf-8") as f:
    json.dump(tables, f, ensure_ascii=False, indent=2, default=str)

print("=== 校验 ===")
print("表1 行数:", len(flows),
      " 总价合计:", round(sum(f["本次总支出"] or 0 for f in flows), 4),
      " 数量合计:", round(sum(f["采购数量"] or 0 for f in flows), 2))
print("表1 采购日期含文字行:", [(i, f["采购日期"], f["备注"]) for i, f in enumerate(flows)
                                  if f["采购日期"] and not is_date_str(f["采购日期"])])
print("表2 SKU数:", len(tables["SKU档案库"]), " 唯一名:", len({s["产品完整名称"] for s in tables["SKU档案库"]}),
      " 总金额合计:", round(sum(s["累计采购总金额"] for s in tables["SKU档案库"]), 4))
print("表3 行数:", len(mat), " 金额合计:", round(sum(m["支出金额"] or 0 for m in mat), 4))
print("表4 行数:", len(recon), " 首行:", recon[0], " 末行:", recon[-1])
print("表5a Y2:", [(y["产品名"], y["毛利率"]) for y in y2])
print("表5b 备货:", len(stock), " 示例:", {k: stock[0][k] for k in ("产品名称", "在途可用库存", "安全采购量(7天)", "建议采购总金额")})
print("图片:", len([s for s in tables['SKU档案库'] if s['_img']]), "张")
