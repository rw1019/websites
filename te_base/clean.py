"""把 Te 数据表.xlsx 清洗成规范的多张表 + 导出内嵌图片，供导入飞书多维表格。"""
import json
import math
import os
import re
import shutil
import xml.etree.ElementTree as ET
from datetime import datetime

import openpyxl
from openpyxl.utils import get_column_letter

SRC = "/workspace/.uploads/04716e07-c586-4875-b818-21b23fe18e64_Te 数据表.xlsx"
XDIR = "/workspace/te_base/xlsx"
OUT = "/workspace/te_base/out"
IMG = "/workspace/te_base/images"

NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def clean_header(s):
    if s is None:
        return None
    return re.sub(r"\s+", "", str(s)).strip()


def to_date(v):
    """把 datetime 或 Excel 序列号转成 'YYYY-MM-DD'；无法转换返回原值。"""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)) and 30000 <= v <= 60000:
        try:
            return openpyxl.utils.datetime.from_excel(v).strftime("%Y-%m-%d")
        except Exception:
            return v
    return v


def sanitize_text(v):
    if isinstance(v, str):
        v = re.sub(r"[\x00-\x08\x0b-\x1f\x7f]", " ", v).strip()
        return v or None
    return v


def is_date_str(v):
    return isinstance(v, str) and bool(re.match(r"^\d{4}-\d{2}-\d{2}$", v))


def to_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except Exception:
        return None


# ---------- 1. 解析内嵌图片映射 ID -> 文件 ----------
rels = ET.parse(f"{XDIR}/xl/_rels/cellimages.xml.rels").getroot()
rid2file = {}
for rel in rels:
    rid2file[rel.get("Id")] = os.path.basename(rel.get("Target"))

cellimg = ET.parse(f"{XDIR}/xl/cellimages.xml").getroot()
id2file = {}
for ci in cellimg:
    pic = ci.find("xdr:pic", NS)
    if pic is None:
        continue
    name = pic.find("xdr:nvPicPr", NS).find("xdr:cNvPr", NS).get("name")  # ID_xxx
    rid = pic.find("xdr:blipFill", NS).find("a:blip", NS).get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
    id2file[name] = rid2file.get(rid)
print(f"内嵌图片映射: {len(id2file)} 张")


def resolve_dispimg(cellval):
    """=DISPIMG("ID_...",1) -> 图片文件名"""
    if not isinstance(cellval, str) or "DISPIMG" not in cellval:
        return None
    m = re.search(r'DISPIMG\("([^"]+)"', cellval)
    return id2file.get(m.group(1)) if m else None


os.makedirs(IMG, exist_ok=True)
os.makedirs(OUT, exist_ok=True)
for f in os.listdir(IMG):
    os.remove(os.path.join(IMG, f))

wb = openpyxl.load_workbook(SRC, data_only=True)
tables = {}


# ---------- 2. 库存：拆成 产品 / 库存流水 ----------
ws = wb["库存"]
hdrs = [clean_header(ws.cell(1, c).value) for c in range(1, 31)]
# 产品块首行 = A 列合并区的左上角
block_rows = sorted(
    {r for rng in ws.merged_cells.ranges if str(rng).startswith("A") for r in [rng.min_row]}
)
block_rows = [r for r in block_rows if isinstance(ws.cell(r, 1).value, str)]
block_rows.sort()

# 区块层级：带缩略图(第 O 列 DISPIMG)的块 = 真产品；紧跟其后无图的块 = 该产品的颜色/规格
bounds = block_rows + [ws.max_row + 1]
blocks = []
for i, hr in enumerate(block_rows):
    blocks.append({
        "hr": hr,
        "end": bounds[i + 1],
        "raw": str(ws.cell(hr, 1).value).strip(),
        "img": resolve_dispimg(ws.cell(hr, 15).value),
    })
parent = None
for b in blocks:
    if b["img"]:  # 有缩略图 = 真产品
        parent = b["raw"]
        b["真产品"], b["款式"] = b["raw"], None
    else:          # 无缩略图 = 上级产品的款式
        b["真产品"], b["款式"] = (parent or b["raw"]), b["raw"]

products, flows = [], []
for b in blocks:
    hr, end = b["hr"], b["end"]
    if b["img"]:
        shutil.copy(os.path.join(XDIR, "xl/media", b["img"]), os.path.join(IMG, f"{b['raw']}_{b['img']}"))
        products.append({
            "产品名": b["真产品"],
            "缩略图": f"{b['raw']}_{b['img']}",
            "SKC": sanitize_text(ws.cell(hr, 16).value),
            "备注": sanitize_text(ws.cell(hr, 17).value),
        })
    for r in range(hr + 1, end):
        if ws.cell(r, 4).value is None and ws.cell(r, 9).value is None:
            continue
        pdate, sdate = to_date(ws.cell(r, 1).value), to_date(ws.cell(r, 2).value)
        notes = []
        n0 = sanitize_text(ws.cell(r, 17).value)
        if n0 is not None:
            notes.append(str(n0))
        for label, d in (("采购", pdate), ("发货", sdate)):
            if d is not None and not is_date_str(d):
                notes.append(f"{label}日期备注：{d}")
                if label == "采购":
                    pdate = None
                else:
                    sdate = None
        flows.append({
            "产品名": b["真产品"],
            "款式": b["款式"],
            "采购日期": pdate,
            "发货日期": sdate,
            "单价": to_num(ws.cell(r, 4).value),
            "运费": to_num(ws.cell(r, 5).value),
            "TEMU运费": to_num(ws.cell(r, 6).value),
            "广告平均单价": to_num(ws.cell(r, 7).value),
            "成本(单个)": to_num(ws.cell(r, 8).value),
            "数量": to_num(ws.cell(r, 9).value),
            "申报价": to_num(ws.cell(r, 10).value),
            "净利率": to_num(ws.cell(r, 11).value),
            "每单赚": to_num(ws.cell(r, 12).value),
            "总价": to_num(ws.cell(r, 13).value),
            "单个重量(kg)": to_num(ws.cell(r, 14).value),
            "备注": "；".join(n for n in notes if n) or None,
        })
tables["产品"] = products
tables["库存流水"] = flows
print(f"库存 -> 产品 {len(products)} 条, 流水 {len(flows)} 条")


def dump_sheet(sheet, hdr_row_range, tname, date_cols=(), num_cols=(), max_col=None):
    ws = wb[sheet]
    mc = max_col or ws.max_column
    hdr = [clean_header(ws.cell(hdr_row_range, c).value) for c in range(1, mc + 1)]
    rows = []
    for r in range(hdr_row_range + 1, ws.max_row + 1):
        if all(ws.cell(r, c).value is None for c in range(1, mc + 1)):
            continue
        rec = {}
        for c in range(1, mc + 1):
            h = hdr[c - 1]
            if not h or isinstance(ws.cell(hdr_row_range, c).value, (int, float)):
                continue  # 跳过表头杂物
            v = ws.cell(r, c).value
            if c in date_cols:
                v = to_date(v)
                if v is not None and not is_date_str(v):
                    v = None
            elif c in num_cols:
                v = to_num(v)
            else:
                v = sanitize_text(v)
            if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
                v = None
            rec[h] = v
        rows.append(rec)
    tables[tname] = rows
    print(f"{sheet} -> {tname}: {len(rows)} 条")


# Y2：利润测算
dump_sheet("Y2", 1, "Y2核算", date_cols=(1,), num_cols=tuple(range(3, 15)))
# 物料：只取 A-C
dump_sheet("物料", 1, "物料采购", date_cols=(1,), num_cols=(3,), max_col=3)
# 春节备货
dump_sheet("春节备货", 1, "春节备货", date_cols=(10,), num_cols=(2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 13, 14))
# 成本：只取 A-B
dump_sheet("成本", 1, "固定成本", num_cols=(2,), max_col=2)
# TK 基本为空，跳过

with open(f"{OUT}/tables.json", "w", encoding="utf-8") as f:
    json.dump(tables, f, ensure_ascii=False, indent=2, default=str)
print("写出:", f"{OUT}/tables.json")
print("图片目录:", IMG, len(os.listdir(IMG)), "个文件")
